from __future__ import annotations

import asyncio
import contextlib
import hashlib
import json
import os
import re
import zipfile
from pathlib import Path, PurePosixPath
from typing import Any, AsyncIterator
from uuid import uuid4

from agents import Runner, set_default_openai_client
from agents.run import RunConfig
from agents.sandbox import Manifest, SandboxAgent, SandboxRunConfig
from agents.sandbox.capabilities import Capabilities
from agents.sandbox.capabilities.skills import Skills
from agents.sandbox.entries import File, LocalDir
from agents.stream_events import AgentUpdatedStreamEvent, RunItemStreamEvent, StreamEvent
from docx import Document
from dotenv import load_dotenv
from openai import AsyncOpenAI
from openpyxl import load_workbook
from pptx import Presentation

from .models import (
    ARTIFACT_FORMAT_EXTENSIONS,
    ARTIFACT_FORMAT_MIME_TYPES,
    FileInputSummary,
    WorkflowDebugTraceEntry,
    WorkflowArtifactRef,
    WorkflowDefinition,
    WorkflowExecutionArtifact,
    WorkflowPersistedResult,
    WorkflowProgressEvent,
    WorkflowRunInputSummary,
    WorkflowRunRecord,
    WorkflowRunTerminalEvent,
    WorkflowRunTimelineEntry,
    WorkflowRunnerEvent,
    utc_now_iso,
    validate_execution_result,
)
from .storage import (
    run_artifact_dir,
    run_upload_dir,
    safe_filename,
    save_run_record,
)
from .workflow_registry import get_workflow

load_dotenv()
SKILLS_SOURCE_DIR = Path("sandflow/sandbox_skills")
DOCKER_SANDBOX_DIR = Path("sandflow/docker_sandbox")
DOCKER_SANDBOX_DOCKERFILE = DOCKER_SANDBOX_DIR / "Dockerfile"
DOCKER_SANDBOX_REQUIREMENTS = DOCKER_SANDBOX_DIR / "requirements.txt"
DOCKER_SANDBOX_IMAGE_PREFIX = "sandflow-sandbox"
SANDBOX_REQUIRED_MODULES = ("docx", "pptx", "openpyxl", "pypdf")
PROGRESS_EVENT_LIMIT = 50
TIMELINE_SUMMARY_LIMIT = 20


def execution_enabled() -> bool:
    return bool(os.getenv("OPENAI_API_KEY") and os.getenv("OPENAI_SANDBOX_MODEL"))


def configure_openai_client() -> None:
    api_key = os.getenv("OPENAI_API_KEY")
    if not api_key:
        return

    base_url = os.getenv("OPENAI_API_BASE")
    kwargs: dict[str, Any] = {"api_key": api_key}
    if base_url:
        kwargs["base_url"] = base_url

    set_default_openai_client(
        AsyncOpenAI(**kwargs),
        use_for_tracing=False,
    )


async def run_workflow(
    workflow_id: str,
    text_inputs: dict[str, str],
    file_inputs: dict[str, Path],
    *,
    debug: bool = False,
) -> WorkflowRunRecord:
    terminal_event: WorkflowRunTerminalEvent | None = None
    async for event in stream_workflow(workflow_id, text_inputs, file_inputs, debug=debug):
        if isinstance(event, WorkflowRunTerminalEvent):
            terminal_event = event

    if terminal_event is None:
        raise RuntimeError("Workflow run ended without a terminal event.")
    if terminal_event.status == "complete" and terminal_event.record is not None:
        return terminal_event.record

    error_message = terminal_event.error
    if not error_message and terminal_event.record is not None:
        error_message = terminal_event.record.error
    raise RuntimeError(error_message or "Workflow run failed.")


async def stream_workflow(
    workflow_id: str,
    text_inputs: dict[str, str],
    file_inputs: dict[str, Path],
    *,
    debug: bool = False,
) -> AsyncIterator[WorkflowRunnerEvent]:
    workflow = get_workflow(workflow_id)
    started_at = utc_now_iso()
    run_id = f"run_{uuid4().hex[:12]}"
    raw_result_json: str | None = None
    timeline_summary: list[WorkflowRunTimelineEntry] = []
    debug_trace: list[WorkflowDebugTraceEntry] = []
    input_summary = WorkflowRunInputSummary()

    def remember(event: WorkflowProgressEvent) -> None:
        if event.persist or event.kind == "error":
            timeline_summary.append(
                WorkflowRunTimelineEntry(
                    timestamp=event.timestamp,
                    stage=event.stage,
                    title=event.title,
                    detail=event.detail,
                )
            )
            if len(timeline_summary) > TIMELINE_SUMMARY_LIMIT:
                del timeline_summary[:-TIMELINE_SUMMARY_LIMIT]

    async def emit(event: WorkflowProgressEvent) -> AsyncIterator[WorkflowRunnerEvent]:
        remember(event)
        yield event

    client, client_options = await _create_sandbox_backend()
    session = None
    try:
        async for event in emit(
            _progress_event(
                stage="preparing",
                kind="stage",
                title="Validating workflow inputs",
                persist=True,
            )
        ):
            yield event
        _validate_run_inputs(workflow, text_inputs, file_inputs)

        async for event in emit(
            _progress_event(
                stage="preparing",
                kind="stage",
                title="Staging uploaded files",
                persist=False,
            )
        ):
            yield event
        input_summary = _stage_run_inputs(run_id, workflow, text_inputs, file_inputs)

        if not execution_enabled():
            raise RuntimeError(
                "Execution is disabled. Set OPENAI_API_KEY and OPENAI_SANDBOX_MODEL before running workflows."
            )

        async for event in emit(
            _progress_event(
                stage="starting_sandbox",
                kind="stage",
                title="Creating sandbox session",
                persist=True,
            )
        ):
            yield event

        configure_openai_client()
        manifest = _build_manifest(workflow, input_summary)
        session = await client.create(manifest=manifest, options=client_options)
        await session.start()
        await _ensure_sandbox_python_runtime(session)

        agent = SandboxAgent(
            name=workflow.name,
            model=os.getenv("OPENAI_SANDBOX_MODEL"),
            instructions=_build_agent_instructions(workflow),
            capabilities=[*Capabilities.default(), _office_artifact_skills()],
        )

        async for event in emit(
            _progress_event(
                stage="running_workflow",
                kind="stage",
                title="Starting workflow agent",
                persist=True,
            )
        ):
            yield event

        run_result = Runner.run_streamed(
            agent,
            "Execute the workflow using the provided workspace inputs and write outputs/result.json.",
            max_turns=50,
            run_config=RunConfig(
                tracing_disabled=True,
                workflow_name=workflow.name,
                sandbox=SandboxRunConfig(client=client, session=session),
            ),
        )

        async for stream_event in run_result.stream_events():
            if debug and _should_capture_debug_trace(stream_event):
                debug_trace.append(_build_debug_trace_entry(stream_event))
            mapped = _map_stream_event_to_progress(stream_event)
            if mapped is not None:
                yield mapped

        async for event in emit(
            _progress_event(
                stage="validating_outputs",
                kind="stage",
                title="Validating outputs/result.json",
                persist=True,
            )
        ):
            yield event

        raw_result_json = await _read_session_text(session, Path("outputs/result.json"))
        validated_result = validate_execution_result(workflow, json.loads(raw_result_json))

        async for event in emit(
            _progress_event(
                stage="saving_outputs",
                kind="stage",
                title="Saving artifacts and run record",
                persist=True,
            )
        ):
            yield event

        artifacts = await _persist_artifacts(run_id, session, validated_result.artifacts)
        record = WorkflowRunRecord(
            id=run_id,
            workflow_id=workflow.id,
            workflow_name=workflow.name,
            workflow_snapshot=workflow,
            status="complete",
            started_at=started_at,
            completed_at=utc_now_iso(),
            input_summary=input_summary,
            result=WorkflowPersistedResult(
                summary=validated_result.summary,
                fields=validated_result.fields,
                artifacts=artifacts,
            ),
            error=None,
            raw_result_json=raw_result_json,
            progress_timeline=list(timeline_summary),
            debug_enabled=debug,
            debug_trace=list(debug_trace),
        )
        save_run_record(record)
        yield WorkflowRunTerminalEvent(status="complete", record=record, error=None)
    except Exception as exc:
        if debug:
            debug_trace.append(
                WorkflowDebugTraceEntry(
                    timestamp=utc_now_iso(),
                    event_type="exception",
                    title="Workflow execution exception",
                    payload=_stringify_debug_value({"error": str(exc)}),
                )
            )
        error_event = _progress_event(
            stage="failed",
            kind="error",
            title="Workflow run failed",
            detail=_truncate_text(str(exc)),
            persist=True,
        )
        remember(error_event)
        failed_record = WorkflowRunRecord(
            id=run_id,
            workflow_id=workflow.id,
            workflow_name=workflow.name,
            workflow_snapshot=workflow,
            status="failed",
            started_at=started_at,
            completed_at=utc_now_iso(),
            input_summary=input_summary,
            result=None,
            error=str(exc),
            raw_result_json=raw_result_json,
            progress_timeline=list(timeline_summary),
            debug_enabled=debug,
            debug_trace=list(debug_trace),
        )
        save_run_record(failed_record)
        yield error_event
        yield WorkflowRunTerminalEvent(
            status="failed",
            record=failed_record,
            error=str(exc),
        )
    finally:
        if session is not None:
            with contextlib.suppress(Exception):
                await session.stop()
            with contextlib.suppress(Exception):
                await client.delete(session)
        docker_sdk_client = getattr(client, "docker_client", None)
        if docker_sdk_client is not None:
            with contextlib.suppress(Exception):
                docker_sdk_client.close()


def _validate_run_inputs(
    workflow: WorkflowDefinition,
    text_inputs: dict[str, str],
    file_inputs: dict[str, Path],
) -> None:
    declared_input_ids = {field.id for field in workflow.input_fields}
    unknown_text_ids = set(text_inputs) - declared_input_ids
    unknown_file_ids = set(file_inputs) - declared_input_ids
    if unknown_text_ids or unknown_file_ids:
        unknown = sorted(unknown_text_ids | unknown_file_ids)
        raise ValueError(f"Unsupported inputs supplied: {', '.join(unknown)}.")

    for field in workflow.input_fields:
        if field.type == "file":
            provided_path = file_inputs.get(field.id)
            if field.required and provided_path is None:
                raise ValueError(f"Missing required file input: {field.label}.")
            if provided_path is not None and not provided_path.exists():
                raise ValueError(f"Uploaded file for `{field.label}` no longer exists.")
            continue

        value = text_inputs.get(field.id, "").strip()
        if field.required and not value:
            raise ValueError(f"Missing required input: {field.label}.")


def _stage_run_inputs(
    run_id: str,
    workflow: WorkflowDefinition,
    text_inputs: dict[str, str],
    file_inputs: dict[str, Path],
) -> WorkflowRunInputSummary:
    text_summary: dict[str, str] = {}
    file_summary: list[FileInputSummary] = []
    upload_dir = run_upload_dir(run_id)

    for field in workflow.input_fields:
        if field.type == "file":
            provided_path = file_inputs.get(field.id)
            if provided_path is None:
                continue
            field_dir = upload_dir / field.id
            field_dir.mkdir(parents=True, exist_ok=True)
            destination = field_dir / safe_filename(provided_path.name)
            destination.write_bytes(provided_path.read_bytes())
            file_summary.append(
                FileInputSummary(
                    input_id=field.id,
                    original_name=provided_path.name,
                    stored_path=str(destination),
                )
            )
            continue

        value = text_inputs.get(field.id, "").strip()
        if value:
            text_summary[field.id] = value

    return WorkflowRunInputSummary(text_fields=text_summary, files=file_summary)


def _build_manifest(workflow: WorkflowDefinition, input_summary: WorkflowRunInputSummary) -> Manifest:
    entries: dict[str, File] = {
        "inputs/workflow_definition.json": File(
            content=workflow.model_dump_json(indent=2).encode("utf-8")
        ),
        "inputs/run_request.json": File(
            content=input_summary.model_dump_json(indent=2).encode("utf-8")
        ),
        "inputs/text_inputs.json": File(
            content=json.dumps(input_summary.text_fields, indent=2).encode("utf-8")
        ),
    }

    for file_input in input_summary.files:
        source = Path(file_input.stored_path)
        entries[f"inputs/files/{file_input.input_id}/{source.name}"] = File(
            content=source.read_bytes()
        )

    return Manifest(entries=entries)


async def _ensure_sandbox_python_runtime(session) -> None:
    check_script = (
        "python3 - <<'PY'\n"
        f"import {', '.join(SANDBOX_REQUIRED_MODULES)}\n"
        "print('sandbox python runtime ready')\n"
        "PY"
    )
    result = await session.exec(check_script, shell=True)
    if result.exit_code == 0:
        return

    stderr_text = result.stderr.decode("utf-8", errors="replace").strip()
    stdout_text = result.stdout.decode("utf-8", errors="replace").strip()
    detail = stderr_text or stdout_text or f"exit code {result.exit_code}"
    raise RuntimeError(
        "Sandbox Python preflight failed. "
        "Office artifact generation requires docx/pptx/openpyxl/pypdf modules. "
        f"Details: {detail}"
    )


async def _create_sandbox_backend():
    try:
        from docker import from_env as docker_from_env
        from docker.errors import BuildError, DockerException, ImageNotFound
        from agents.sandbox.sandboxes.docker import DockerSandboxClient, DockerSandboxClientOptions
    except Exception as exc:
        raise RuntimeError(
            "Docker sandbox support is unavailable. Run `uv sync` so the `docker` package is installed."
        ) from exc

    try:
        docker_client = await asyncio.to_thread(docker_from_env)
        await asyncio.to_thread(docker_client.ping)
    except DockerException as exc:
        raise RuntimeError(
            "Docker sandbox is unavailable. Start Docker and ensure the local daemon is reachable."
        ) from exc

    image_name, managed_image = _sandbox_image_reference()
    if managed_image:
        try:
            await asyncio.to_thread(docker_client.images.get, image_name)
        except ImageNotFound:
            try:
                await asyncio.to_thread(
                    docker_client.images.build,
                    path=str(DOCKER_SANDBOX_DIR.resolve()),
                    dockerfile="Dockerfile",
                    tag=image_name,
                    rm=True,
                    pull=True,
                )
            except (BuildError, DockerException) as exc:
                raise RuntimeError(
                    f"Failed to build Docker sandbox image `{image_name}` from `{DOCKER_SANDBOX_DIR}`."
                ) from exc

    return DockerSandboxClient(docker_client), DockerSandboxClientOptions(image=image_name)


def _sandbox_image_reference() -> tuple[str, bool]:
    configured_image = os.getenv("MONROVIA_SANDBOX_IMAGE", "").strip()
    if configured_image:
        return configured_image, False

    digest = hashlib.sha256()
    for path in (DOCKER_SANDBOX_DOCKERFILE, DOCKER_SANDBOX_REQUIREMENTS):
        digest.update(path.read_bytes())
    return f"{DOCKER_SANDBOX_IMAGE_PREFIX}:{digest.hexdigest()[:12]}", True


def _build_agent_instructions(workflow: WorkflowDefinition) -> str:
    output_schema = {
        "summary": "short run summary shown to the user",
        "fields": {
            field.id: {
                "type": field.type,
                "label": field.label,
                "required": field.required,
                "help_text": field.help_text,
            }
            for field in workflow.output_fields
        },
        "artifacts": [
            {
                "artifact_id": artifact.id,
                "label": artifact.label,
                "format": artifact.format,
                "required": artifact.required,
                "help_text": artifact.help_text,
            }
            for artifact in workflow.artifact_outputs
        ],
    }

    return (
        "You are executing a user-defined workflow inside a sandboxed workspace.\n\n"
        "Read these workspace files before acting:\n"
        "- inputs/workflow_definition.json\n"
        "- inputs/run_request.json\n"
        "- inputs/text_inputs.json\n"
        "- any files under inputs/files/\n\n"
        "Builder-authored workflow prompt:\n"
        f"{workflow.prompt}\n\n"
        "Execution contract:\n"
        "1. Produce the final machine-readable result at outputs/result.json.\n"
        "2. The JSON must include exactly the declared output fields and artifact references.\n"
        "3. Any generated files must be written under outputs/artifacts/.\n"
        "4. Artifact paths in outputs/result.json must be relative workspace paths under outputs/artifacts/.\n"
        "5. Do not fabricate missing inputs. Work only from workspace contents.\n\n"
        "6. For artifact generation, consult the mounted `office-artifacts` skill under `.agents/`.\n"
        "   Use it to learn which libraries are available and how to use them correctly.\n"
        "   Prefer writing Python directly with the appropriate library for the target format.\n"
        "   For example: `python-docx` for docx, `python-pptx` for pptx, `openpyxl` for xlsx, and `csv` for csv.\n"
        "   Choose the document or file structure based on the task itself instead of forcing a generic template.\n"
        "   The helper scripts in `.agents/office-artifacts/scripts/` are optional fallbacks, not the default required path.\n"
        "   Use `python3` to run your code in this environment.\n"
        "   Avoid generic report scaffolding unless the task clearly calls for it.\n"
        "   Do not write placeholder or fallback text into Office files. If generation fails, stop and let the run fail.\n\n"
        "Expected result schema:\n"
        f"{json.dumps(output_schema, indent=2)}\n"
    )


def _progress_event(
    *,
    stage: str,
    kind: str,
    title: str,
    detail: str = "",
    persist: bool = False,
) -> WorkflowProgressEvent:
    return WorkflowProgressEvent(
        timestamp=utc_now_iso(),
        stage=stage,
        kind=kind,
        title=title,
        detail=detail,
        persist=persist,
    )


def _map_stream_event_to_progress(event: StreamEvent) -> WorkflowProgressEvent | None:
    if getattr(event, "type", "") == "raw_response_event":
        return None

    if isinstance(event, AgentUpdatedStreamEvent):
        agent_name = _truncate_text(getattr(event.new_agent, "name", "") or "")
        title = "Agent handoff" if agent_name else "Agent updated"
        return _progress_event(
            stage="running_workflow",
            kind="agent",
            title=title,
            detail=agent_name,
        )

    if not isinstance(event, RunItemStreamEvent):
        return None

    if event.name == "tool_called":
        tool_name = _friendly_tool_name(event.item)
        detail = _truncate_text(_extract_tool_call_detail(event.item))
        return _progress_event(
            stage="running_workflow",
            kind="tool_called",
            title=f"{tool_name} called" if tool_name != "Tool" else "Tool called",
            detail=detail,
        )

    if event.name == "tool_output":
        tool_name = _friendly_tool_name(event.item)
        detail = _truncate_text(_extract_tool_output_detail(event.item))
        return _progress_event(
            stage="running_workflow",
            kind="tool_output",
            title=f"{tool_name} completed" if tool_name != "Tool" else "Tool output received",
            detail=detail,
        )

    if event.name == "message_output_created":
        detail = _truncate_text(_extract_message_detail(event.item))
        return _progress_event(
            stage="running_workflow",
            kind="message",
            title="Agent produced output",
            detail=detail,
        )

    return None


def _build_debug_trace_entry(event: StreamEvent) -> WorkflowDebugTraceEntry:
    event_name = getattr(event, "name", "")
    class_name = event.__class__.__name__
    raw_type = getattr(event, "type", "") or ""
    event_type = event_name or _normalize_event_type(str(raw_type)) or _normalize_event_type(class_name)
    title = class_name if not event_name else f"{class_name}: {event_name}"
    return WorkflowDebugTraceEntry(
        timestamp=utc_now_iso(),
        event_type=str(event_type),
        title=title,
        payload=_stringify_debug_value(_debug_event_payload(event)),
    )


def _should_capture_debug_trace(event: StreamEvent) -> bool:
    return getattr(event, "type", "") != "raw_response_event"


def _normalize_event_type(class_name: str) -> str:
    snake_case = re.sub(r"(?<!^)(?=[A-Z])", "_", class_name).replace("-", "_").lower()
    return snake_case.removesuffix("_stream_event")


async def _read_session_text(session, path: Path) -> str:
    handle = await session.read(path)
    data = handle.read()
    if isinstance(data, bytes):
        return data.decode("utf-8")
    return str(data)


async def _persist_artifacts(
    run_id: str,
    session,
    artifacts: list[WorkflowExecutionArtifact],
) -> list[WorkflowArtifactRef]:
    local_artifact_dir = run_artifact_dir(run_id)
    persisted: list[WorkflowArtifactRef] = []

    for artifact in artifacts:
        relative_path = PurePosixPath(artifact.path)
        filename = relative_path.name
        local_path = local_artifact_dir / safe_filename(filename)
        handle = await session.read(Path(artifact.path))
        payload = handle.read()
        if not isinstance(payload, bytes):
            payload = bytes(payload)
        local_path.write_bytes(payload)
        artifact_format = artifact.format or _format_from_filename(local_path.name)
        _validate_persisted_artifact(local_path, artifact_format)
        persisted.append(
            WorkflowArtifactRef(
                artifact_id=artifact.artifact_id,
                label=artifact.label,
                format=artifact_format,
                stored_path=str(local_path),
                filename=local_path.name,
                mime_type=artifact.mime_type or _mime_type_from_filename(local_path.name),
            )
        )

    return persisted


def _friendly_tool_name(item: Any) -> str:
    title = _safe_attr(item, "title")
    if isinstance(title, str) and title.strip():
        return _truncate_text(title.strip(), limit=40)

    raw_item = _safe_attr(item, "raw_item")
    name = _coerce_lookup(raw_item, "name")
    if isinstance(name, str) and name.strip():
        return _truncate_text(name.strip(), limit=40)

    raw_type = str(_coerce_lookup(raw_item, "type") or "").lower()
    if "shell" in raw_type:
        return "Shell tool"
    if "file_search" in raw_type:
        return "File search"
    if "web_search" in raw_type:
        return "Web search"
    if "computer" in raw_type:
        return "Computer tool"
    if "tool_search" in raw_type:
        return "Tool search"
    return "Tool"


def _extract_tool_call_detail(item: Any) -> str:
    description = _safe_attr(item, "description")
    if isinstance(description, str) and description.strip():
        return description.strip()

    raw_item = _safe_attr(item, "raw_item")
    arguments = _coerce_lookup(raw_item, "arguments")
    if arguments:
        return _stringify_preview(arguments)
    return ""


def _extract_tool_output_detail(item: Any) -> str:
    output = _safe_attr(item, "output")
    if output:
        return _stringify_preview(output)

    raw_item = _safe_attr(item, "raw_item")
    if raw_item:
        text = _coerce_lookup(raw_item, "output_text")
        if text:
            return _stringify_preview(text)
        return _stringify_preview(raw_item)
    return ""


def _extract_message_detail(item: Any) -> str:
    raw_item = _safe_attr(item, "raw_item")
    content = _coerce_lookup(raw_item, "content")
    if isinstance(content, list):
        parts: list[str] = []
        for entry in content:
            text = _coerce_lookup(entry, "text")
            if isinstance(text, str) and text.strip():
                parts.append(text.strip())
                continue
            nested_text = _coerce_lookup(text, "value")
            if isinstance(nested_text, str) and nested_text.strip():
                parts.append(nested_text.strip())
        if parts:
            return " ".join(parts)
    return ""


def _coerce_lookup(value: Any, key: str) -> Any:
    if isinstance(value, dict):
        return value.get(key)
    return getattr(value, key, None)


def _safe_attr(value: Any, name: str) -> Any:
    try:
        return getattr(value, name)
    except Exception:
        return None


def _stringify_preview(value: Any) -> str:
    if isinstance(value, str):
        return value
    if isinstance(value, bytes):
        return value.decode("utf-8", errors="ignore")
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, ensure_ascii=True)
    return str(value)


def _stringify_debug_value(value: Any) -> str:
    if isinstance(value, str):
        return value
    return json.dumps(value, ensure_ascii=True, indent=2)


def _debug_event_payload(value: Any, *, depth: int = 0) -> Any:
    if isinstance(value, AgentUpdatedStreamEvent):
        return {
            "agent_name": getattr(getattr(value, "new_agent", None), "name", "") or "",
        }
    if isinstance(value, RunItemStreamEvent):
        item = getattr(value, "item", None)
        payload = {
            "event_name": value.name,
        }
        if value.name in {"tool_called", "tool_output"}:
            payload["tool_name"] = _friendly_tool_name(item)
        detail = ""
        if value.name == "tool_called":
            detail = _extract_tool_call_detail(item)
        elif value.name == "tool_output":
            detail = _extract_tool_output_detail(item)
        elif value.name == "message_output_created":
            detail = _extract_message_detail(item)
        if detail:
            payload["detail"] = _truncate_text(detail, limit=500)
        return payload
    if depth >= 6:
        return repr(value)
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, bytes):
        return value.decode("utf-8", errors="ignore")
    if isinstance(value, Path):
        return str(value)
    if isinstance(value, dict):
        return {str(key): _debug_event_payload(item, depth=depth + 1) for key, item in value.items()}
    if isinstance(value, (list, tuple, set)):
        return [_debug_event_payload(item, depth=depth + 1) for item in value]
    if hasattr(value, "model_dump"):
        try:
            dumped = value.model_dump(mode="json")
        except TypeError:
            dumped = value.model_dump()
        return _debug_event_payload(dumped, depth=depth + 1)
    if hasattr(value, "__dict__"):
        return {
            key: _debug_event_payload(item, depth=depth + 1)
            for key, item in vars(value).items()
            if not key.startswith("_")
        }
    return repr(value)


def _truncate_text(value: str, *, limit: int = 160) -> str:
    collapsed = " ".join(str(value).split())
    if len(collapsed) <= limit:
        return collapsed
    return f"{collapsed[: limit - 1].rstrip()}…"


def _validate_persisted_artifact(path: Path, artifact_format: str) -> None:
    if artifact_format == "pptx":
        _validate_pptx_file(path)
        return
    if artifact_format == "docx":
        _validate_docx_file(path)
        return
    if artifact_format == "xlsx":
        _validate_xlsx_file(path)
        return


def _validate_pptx_file(path: Path) -> None:
    _validate_zip_entries(path, required_entries={"[Content_Types].xml", "_rels/.rels", "ppt/presentation.xml"})
    try:
        Presentation(str(path))
    except Exception as exc:
        raise ValueError(f"Generated PPTX artifact `{path.name}` is not a valid PowerPoint file.") from exc


def _validate_docx_file(path: Path) -> None:
    _validate_zip_entries(path, required_entries={"[Content_Types].xml", "_rels/.rels", "word/document.xml"})
    try:
        Document(str(path))
    except Exception as exc:
        raise ValueError(f"Generated DOCX artifact `{path.name}` is not a valid Word document.") from exc


def _validate_xlsx_file(path: Path) -> None:
    _validate_zip_entries(path, required_entries={"[Content_Types].xml", "_rels/.rels", "xl/workbook.xml"})
    try:
        workbook = load_workbook(filename=path, read_only=True)
        workbook.close()
    except Exception as exc:
        raise ValueError(f"Generated XLSX artifact `{path.name}` is not a valid Excel workbook.") from exc


def _validate_zip_entries(path: Path, *, required_entries: set[str]) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            members = set(archive.namelist())
    except Exception as exc:
        raise ValueError(f"Generated artifact `{path.name}` is not a valid Office package.") from exc

    missing = sorted(required_entries - members)
    if missing:
        raise ValueError(
            f"Generated artifact `{path.name}` is missing required Office package entries: {', '.join(missing)}."
        )


__all__ = ["execution_enabled", "run_workflow", "stream_workflow"]


def _office_artifact_skills() -> Skills:
    return Skills(from_=LocalDir(src=SKILLS_SOURCE_DIR))


def _format_from_filename(filename: str) -> str:
    suffix = Path(filename).suffix.lower()
    for format_name, extension in ARTIFACT_FORMAT_EXTENSIONS.items():
        if suffix == extension:
            return format_name
    return "txt"


def _mime_type_from_filename(filename: str) -> str | None:
    return ARTIFACT_FORMAT_MIME_TYPES.get(_format_from_filename(filename))
